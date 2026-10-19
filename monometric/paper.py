"""Publication outputs for an ISPRS full paper.

    monometric summarize <cartella_report> -o paper/      -> CSV + tables.docx + tables.tex
    monometric figures   <cartella_report> -o paper/fig/  -> PDF (vector) + PNG 300 dpi

Layout from the ISPRS full-paper guidelines: A4, two columns 82 mm wide with 6 mm spacing
(full text width 170 mm), Times New Roman. Figure text is set at 8 pt so that it reads like
the 9 pt body text. Captions are NOT drawn in the figures: ISPRS wants them centred under
the illustration in the manuscript (suggested captions are written to captions.txt).
Labels are in English (language of the paper). Nothing identifies the authors (double-blind).
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

import numpy as np

COL_MM, PAGE_MM = 82.0, 170.0
MM = 1 / 25.4
SERIES = ["#2a78d6", "#eb6834", "#1baf7a"]          # categorical slots 1-3 (validated, all pairs)
MARKERS = ["o", "s", "^"]                            # secondary encoding (print / CVD)
INK, INK2, GRID = "#0b0b0b", "#52514e", "#d9d8d3"
SCALES = {"1:50": 10.0, "1:100": 20.0, "1:200": 40.0, "1:500": 100.0}   # 0.2 mm graphic error x scale [mm]


# =============================================================================== loading
def load_reports(root):
    root = Path(root)
    paths = sorted(root.rglob("*_report.json")) if root.is_dir() else [root]
    reps = []
    for p in paths:
        try:
            r = json.loads(p.read_text(encoding="utf-8"))
        except Exception:
            continue
        r["_path"] = str(p)
        r["_dir"] = str(p.parent)
        r["_stem"] = p.name[: -len("_report.json")]
        reps.append(r)
    if not reps:
        raise RuntimeError(f"Nessun *_report.json trovato in {root}")
    return reps


def _device(r):
    ex = r.get("image", {}).get("exif", {})
    return f"{ex.get('Make', '')} {ex.get('Model', '')}".strip() or "unknown"


def _calib_short(r):
    src = r.get("camera", {}).get("source", "")
    return "calibrated" if src.startswith("calibrazione") else "user K" if "utente" in src else "single image"


def condition(r):
    if r.get("tag"):
        return r["tag"]
    fam = r["models"][r["chosen"]]["family"]
    mode = "zero-shot" if fam == "none" else "anchored"
    return f"{r['chosen']} {mode} ({_calib_short(r)})"


def _chosen_per_board(r):
    m = r["models"][r["chosen"]]
    v = (m.get("cv") or {}).get(m["family"]) or {}
    return v.get("per_board") or []


def _val_errors(r):
    """Per-point validation errors [mm] and their source: check points if present, else LOBO board RMSE."""
    g = r.get("gcp") or {}
    if g.get("checks"):
        return np.array([c["err_mm"] for c in g["checks"]]), "check points"
    pb = _chosen_per_board(r)
    return np.array([b["rmse3d_mm"] for b in pb]), "boards (LOBO)"


# =============================================================================== summarize
def summarize(root, out):
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    reps = load_reports(root)
    runs, boards, checks, sweeps, conds = [], [], [], [], []
    for r in reps:
        cam, cv, g = r["camera"], r.get("chosen_cv") or {}, r.get("gcp") or {}
        base = dict(site=r.get("site") or Path(r["_dir"]).name, image=r["_stem"], condition=condition(r),
                    device=_device(r), model=r["chosen"], family=r["models"][r["chosen"]]["family"],
                    camera=_calib_short(r))
        runs.append(dict(base, width=r["image"]["width"], height=r["image"]["height"], f_px=cam["fx"],
                         sigma_f_pct=100 * cam["sigma_f_rel"], distortion=cam["distortion_model"],
                         n_boards=len(r["detections"]), n_corners=sum(d["n_corners"] for d in r["detections"]),
                         reproj_rms_px=cam["rms_px"],
                         lobo_rmse3d_mm=cv.get("rmse3d_mm"), lobo_mae3d_mm=cv.get("mae3d_mm"), lobo_d90_mm=cv.get("d90_mm"),
                         lobo_radial_rmse_mm=cv.get("radial_rmse_mm"), lobo_lateral_rmse_mm=cv.get("lateral_rmse_mm"),
                         lobo_rel_depth_err_pct=cv.get("rel_depth_err_pct"), dist_scale_err_pct=cv.get("dist_scale_err_pct"),
                         sigma_point_pct=100 * (r.get("sigma_rel_point") or 0),
                         control_rmse_mm=g.get("control_rmse_mm"), helmert_scale_ppm=g.get("helmert_scale_ppm"),
                         n_check=g.get("n_check"), check_mae_mm=g.get("check_mae_mm"), check_rmse_mm=g.get("check_rmse_mm"),
                         check_d90_mm=g.get("check_d90_mm"), check_max_mm=g.get("check_max_mm"),
                         check_radial_rmse_mm=g.get("check_radial_rmse_mm"),
                         check_lateral_rmse_mm=g.get("check_lateral_rmse_mm"),
                         coverage_1s=g.get("coverage_1s"), coverage_2s=g.get("coverage_2s"),
                         n_points=r.get("n_points"), runtime_s=r.get("runtime_s"),
                         profile=(r.get("hardware") or {}).get("name")))
        for b in _chosen_per_board(r):
            boards.append(dict(base, **b))
        for c in g.get("checks") or []:
            checks.append(dict(base, **{k: v for k, v in c.items() if k != "dxyz_mm"},
                               dx_mm=c["dxyz_mm"][0], dy_mm=c["dxyz_mm"][1], dz_mm=c["dxyz_mm"][2]))
        for s in r.get("focal_sweep") or []:
            sweeps.append(dict(base, **s))
        for mk, m in r["models"].items():
            for fam, v in (m.get("cv") or {}).items():
                s = (v or {}).get("summary") or {}
                if s:
                    conds.append(dict(base, candidate_model=mk, candidate_family=fam, chosen=(mk == r["chosen"] and fam == m["family"]), **s))
    files = {}
    for name, rows in [("runs", runs), ("boards", boards), ("checks", checks), ("focal_sweep", sweeps),
                       ("model_family_lobo", conds)]:
        if rows:
            files[name] = _write_csv(out / f"{name}.csv", rows)
    tables = _tables(runs, checks, boards)
    files["tables_tex"] = _write_tex(out / "tables.tex", tables)
    try:
        files["tables_docx"] = _write_docx(out / "tables.docx", tables)
    except ImportError:
        files["tables_docx"] = "skipped (pip install python-docx)"
    return files


def _write_csv(path, rows):
    keys = []
    for r in rows:
        for k in r:
            if k not in keys:
                keys.append(k)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=keys)
        w.writeheader()
        for r in rows:
            w.writerow({k: (round(v, 4) if isinstance(v, float) else v) for k, v in r.items()})
    return str(path)


def _f(v, nd=1):
    return "–" if v is None or (isinstance(v, float) and not np.isfinite(v)) else f"{v:.{nd}f}"


def _tables(runs, checks, boards):
    T = []
    # Table A - accuracy per site and condition
    groups = {}
    for r in runs:
        groups.setdefault((r["site"], r["condition"]), []).append(r)
    rowsA = []
    for (site, cond), rs in sorted(groups.items()):
        ck = [c for c in checks if c["site"] == site and c["condition"] == cond]
        if ck:
            e = np.array([c["err_mm"] for c in ck])
            rad = np.array([c["radial_mm"] for c in ck])
            lat = np.array([c["lateral_mm"] for c in ck])
            src = "CP"
        else:
            bd = [b for b in boards if b["site"] == site and b["condition"] == cond]
            e = np.array([b["rmse3d_mm"] for b in bd]) if bd else np.array([])
            rad = np.array([b["radial_rmse_mm"] for b in bd]) if bd else np.array([])
            lat = np.array([b["lateral_rmse_mm"] for b in bd]) if bd else np.array([])
            src = "LOBO"
        rowsA.append([site, cond, str(len(rs)), f"{len(e)} {src}",
                      _f(e.mean() if e.size else None), _f(np.sqrt(np.mean(e ** 2)) if e.size else None),
                      _f(np.percentile(e, 90) if e.size else None),
                      _f(np.sqrt(np.mean(rad ** 2)) if rad.size else None), _f(np.sqrt(np.mean(lat ** 2)) if lat.size else None)])
    T.append(dict(title="Accuracy per site and processing condition (CP = independent check points, LOBO = "
                        "leave-one-board-out). Errors in mm.",
                  head=["Site", "Condition", "Images", "Points", "MAE", "RMSE", "D90", "Radial RMSE", "Lateral RMSE"],
                  rows=rowsA))
    # Table B - camera / focal
    rowsB = []
    seen = set()
    for r in runs:
        k = (r["device"], r["camera"], round(r["f_px"], 1))
        if k in seen:
            continue
        seen.add(k)
        rowsB.append([r["device"], r["camera"], f"{r['width']}×{r['height']}", _f(r["f_px"], 1),
                      _f(r["sigma_f_pct"], 2), r["distortion"], _f(r["reproj_rms_px"], 2)])
    T.append(dict(title="Camera parameters used (σf = relative 1-sigma of the focal length).",
                  head=["Device", "Intrinsics", "Image (px)", "f (px)", "σf (%)", "Distortion", "Reproj. RMS (px)"],
                  rows=rowsB))
    # Table C - representation scale compliance
    rowsC = []
    for (site, cond), rs in sorted(groups.items()):
        ck = [c["err_mm"] for c in checks if c["site"] == site and c["condition"] == cond]
        if not ck:
            ck = [b["rmse3d_mm"] for b in boards if b["site"] == site and b["condition"] == cond]
        if not ck:
            continue
        e = np.array(ck)
        rowsC.append([site, cond] + [f"{100 * np.mean(e <= tol):.0f}" for tol in SCALES.values()])
    T.append(dict(title="Share of validation points within the graphic tolerance of each representation scale "
                        "(0.2 mm × scale denominator), in %.",
                  head=["Site", "Condition"] + [f"{k} (≤{v:.0f} mm)" for k, v in SCALES.items()], rows=rowsC))
    return T


def _write_tex(path, tables):
    out = []
    for i, t in enumerate(tables, 1):
        n = len(t["head"])
        out.append("\\begin{table*}[ht]\n\\centering\\footnotesize\n\\begin{tabular}{" + "l" * 2 + "r" * (n - 2) + "}\n\\hline")
        out.append(" & ".join(t["head"]) + " \\\\\n\\hline")
        for r in t["rows"]:
            out.append(" & ".join(str(x).replace("%", "\\%").replace("_", "\\_") for x in r) + " \\\\")
        out.append("\\hline\n\\end{tabular}\n\\caption{" + t["title"].replace("%", "\\%") + "}\n\\label{tab:mm" + str(i) +
                   "}\n\\end{table*}\n")
    Path(path).write_text("\n".join(out), encoding="utf-8")
    return str(path)


def _write_docx(path, tables):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(9)
    for i, t in enumerate(tables, 1):
        tab = doc.add_table(rows=1, cols=len(t["head"]))
        tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        tab.style = "Table Grid"
        for c, h in zip(tab.rows[0].cells, t["head"]):
            c.text = h
            for p in c.paragraphs:
                for run in p.runs:
                    run.bold = True
        for r in t["rows"]:
            cells = tab.add_row().cells
            for c, v in zip(cells, r):
                c.text = str(v)
        for row in tab.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.size, run.font.name = Pt(8), "Times New Roman"
        cap = doc.add_paragraph(f"Table {i}. {t['title']}")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()
    doc.save(str(path))
    return str(path)


# =============================================================================== figures
def _style():
    import matplotlib as mpl
    from matplotlib import font_manager

    names = {f.name for f in font_manager.fontManager.ttflist}
    serif = [n for n in ("Times New Roman", "Liberation Serif", "Nimbus Roman", "TeX Gyre Termes", "DejaVu Serif") if n in names]
    mpl.rcParams.update({
        "font.family": "serif", "font.serif": serif or ["DejaVu Serif"], "font.size": 8,
        "axes.labelsize": 8, "axes.titlesize": 8, "xtick.labelsize": 7, "ytick.labelsize": 7, "legend.fontsize": 7,
        "axes.edgecolor": INK2, "axes.labelcolor": INK, "xtick.color": INK2, "ytick.color": INK2,
        "axes.linewidth": 0.6, "xtick.major.width": 0.6, "ytick.major.width": 0.6,
        "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.4, "axes.axisbelow": True,
        "axes.spines.top": False, "axes.spines.right": False, "legend.frameon": False,
        "lines.linewidth": 1.2, "lines.markersize": 4, "savefig.dpi": 300, "pdf.fonttype": 42, "ps.fonttype": 42,
        "figure.dpi": 150,
    })
    return serif[0] if serif else "DejaVu Serif"


def _save(fig, out, name, captions, caption):
    # exact physical size (82 mm column / 170 mm page): insert at 100 % and text stays at 8 pt
    fig.savefig(out / f"{name}.pdf", metadata={"Creator": None, "Author": None})
    fig.savefig(out / f"{name}.png", dpi=300)
    captions.append(f"{name}: {caption}")
    import matplotlib.pyplot as plt

    plt.close(fig)


def figures(root, out, run=None, cond=None):
    """cond: condition shown in the single-condition figures (error vs range, coverage, orientation);
    default = the condition with the lowest median validation error."""
    import matplotlib.pyplot as plt

    font = _style()
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    allreps = load_reports(root)
    caps, made = [], []
    errs = {}
    for r in allreps:
        errs.setdefault(condition(r), []).extend(_val_errors(r)[0].tolist())
    if cond is None:
        cond = min(errs, key=lambda c: np.median(errs[c]) if errs[c] else np.inf)
    reps = [r for r in allreps if condition(r) == cond] or allreps

    # ---- 1. qualitative panel for one run (full page width)
    sel = [r for r in reps if run is None or r["_stem"] == run] or reps[:1]
    for r in sel[:1] if run is None else sel:
        p = _qualitative(r, out, caps, plt)
        if p:
            made.append(p)

    # ---- 2. error vs range, radial / lateral
    pts = []
    for r in reps:
        g = r.get("gcp") or {}
        if g.get("checks"):
            pts += [(c["range_m"], abs(c["radial_mm"]), c["lateral_mm"]) for c in g["checks"]]
    src = "independent check points"
    if not pts:
        src = "ChArUco boards, leave-one-board-out"
        for r in reps:
            pts += [(b["range_m"], b["radial_rmse_mm"], b["lateral_rmse_mm"]) for b in _chosen_per_board(r) if "range_m" in b]
    if pts:
        P = np.array(pts)
        fig, ax = plt.subplots(layout="constrained", figsize=(COL_MM * MM, 55 * MM))
        ax.scatter(P[:, 0], P[:, 1], s=14, marker=MARKERS[0], color=SERIES[0], edgecolors="white", linewidths=0.4,
                   label="radial (along the ray)", zorder=3)
        ax.scatter(P[:, 0], P[:, 2], s=14, marker=MARKERS[1], color=SERIES[1], edgecolors="white", linewidths=0.4,
                   label="lateral", zorder=3)
        ax.set_xlabel("Range from camera (m)")
        ax.set_ylabel("Error (mm)")
        ax.set_ylim(bottom=0)
        ax.legend(loc="upper left")
        _save(fig, out, "fig_error_vs_range", caps,
              f"Radial and lateral error as a function of camera range ({src}; condition: {cond}).")
        made.append("fig_error_vs_range")

    # ---- 3. focal sensitivity
    sw = [dict(s, condition=condition(r)) for r in allreps for s in (r.get("focal_sweep") or [])]
    if sw:
        cn = list(dict.fromkeys(s["condition"] for s in sw))[:3]
        eps = sorted({s["eps"] for s in sw})
        E = np.array(eps) * 100
        lim = max(abs(E).max(), 1) * 1.15
        fig, axs = plt.subplots(2, 1, layout="constrained", sharex=True, figsize=(COL_MM * MM, 95 * MM))
        for ax, key, ttl in [(axs[0], "axial_rel_median", "(a) along the optical axis, ΔZ/Z"),
                             (axs[1], "transverse_rel_median", "(b) in the image plane, |ΔX, ΔY|/Z")]:
            if key == "axial_rel_median":
                ax.plot([-lim, lim], [-lim, lim], ls="--", lw=0.8, color=INK2, zorder=1, label="first-order theory")
            else:
                ax.axhline(0, ls="--", lw=0.8, color=INK2, zorder=1, label="first-order theory")
            for k, c in enumerate(cn):
                y = [np.mean([s[key] for s in sw if s["eps"] == e and s["condition"] == c] or [np.nan]) * 100 for e in eps]
                ax.plot(E, y, marker=MARKERS[k], color=SERIES[k], label=c, zorder=3)
            ax.set_title(ttl, loc="left")
            ax.set_ylabel("Median displacement (%)")
        axs[1].set_xlabel("Relative focal-length error ε (%)")
        axs[0].legend(loc="upper left", fontsize=6.5)
        _save(fig, out, "fig_focal_sensitivity", caps,
              "Sensitivity of the reconstruction to a relative error ε of the focal length: the whole chain "
              "(PnP anchors, network conditioning, alignment) is re-run with f(1+ε) and compared with the "
              "nominal solution. Dashed lines: first-order prediction ΔZ/Z = ε along the optical axis, "
              "X and Y unchanged.")
        made.append("fig_focal_sensitivity")
        if any(s.get("check_rmse_mm") is not None for s in sw):
            fig, axs = plt.subplots(1, 2, layout="constrained", sharey=True, figsize=(COL_MM * MM, 50 * MM))
            for ax, key, ttl in [(axs[0], "check_radial_rmse_mm", "(a) radial"), (axs[1], "check_lateral_rmse_mm", "(b) lateral")]:
                for k, c in enumerate(cn):
                    y = [np.mean([s[key] for s in sw if s["eps"] == e and s["condition"] == c and s.get(key) is not None]
                                 or [np.nan]) for e in eps]
                    ax.plot(E, y, marker=MARKERS[k], color=SERIES[k], label=c)
                ax.set_title(ttl, loc="left")
                ax.set_xlabel("ε (%)")
            axs[0].set_ylabel("Check-point RMSE (mm)")
            axs[0].set_ylim(bottom=0)
            axs[1].legend(loc="upper center", fontsize=6)
            _save(fig, out, "fig_focal_sensitivity_checkpoints", caps,
                  "Check-point RMSE, radial and lateral components, as a function of the focal-length error ε.")
            made.append("fig_focal_sensitivity_checkpoints")

    # ---- 4. conditions comparison
    conds = errs
    if len(conds) >= 2:
        names = list(conds)
        fig, ax = plt.subplots(layout="constrained", figsize=((COL_MM if len(names) <= 4 else PAGE_MM) * MM, 60 * MM))
        bp = ax.boxplot([conds[n] for n in names], widths=0.5, patch_artist=True, showfliers=True,
                        medianprops=dict(color=INK, lw=1.0), whiskerprops=dict(color=INK2, lw=0.6),
                        capprops=dict(color=INK2, lw=0.6), flierprops=dict(marker="o", ms=2.5, mec=INK2, mfc="none"))
        for b in bp["boxes"]:
            b.set(facecolor=SERIES[0] + "33", edgecolor=SERIES[0], lw=0.8)
        ax.set_xticks(range(1, len(names) + 1))
        ax.set_xticklabels([n.replace(" (", "\n(") for n in names], rotation=0)
        ax.set_ylabel("3D error (mm)")
        ax.set_ylim(bottom=0)
        ax.grid(axis="x", visible=False)
        _save(fig, out, "fig_conditions", caps,
              "Distribution of the 3D error for each processing condition (box: interquartile range, line: median).")
        made.append("fig_conditions")

    # ---- 5. uncertainty coverage
    cv = [(c["sigma_pred_mm"], abs(c["radial_mm"])) for r in reps for c in ((r.get("gcp") or {}).get("checks") or [])]
    if cv:
        C = np.array(cv)
        fig, ax = plt.subplots(layout="constrained", figsize=(COL_MM * MM, 60 * MM))
        m = max(C.max() * 1.1, 1)
        ax.plot([0, m], [0, m], ls="--", lw=0.8, color=INK2)
        ax.plot([0, m / 2], [0, m], ls=":", lw=0.8, color=INK2)
        ax.text(m * 0.98, m * 0.93, "1σ", ha="right", fontsize=7, color=INK2)
        ax.text(m * 0.47, m * 0.95, "2σ", ha="right", fontsize=7, color=INK2)
        ax.scatter(C[:, 0], C[:, 1], s=14, marker=MARKERS[0], color=SERIES[0], edgecolors="white", linewidths=0.4, zorder=3)
        ax.set_xlim(0, m)
        ax.set_ylim(0, m)
        ax.set_aspect("equal")
        ax.set_xlabel("Predicted σ (mm)")
        ax.set_ylabel("|Radial error| at check point (mm)")
        c1, c2 = np.mean(C[:, 1] <= C[:, 0]), np.mean(C[:, 1] <= 2 * C[:, 0])
        ax.text(0.03 * m, 0.95 * m, f"within 1σ: {100 * c1:.0f}%\nwithin 2σ: {100 * c2:.0f}%", va="top", fontsize=7, color=INK)
        _save(fig, out, "fig_uncertainty_coverage", caps,
              "Declared per-point uncertainty versus observed radial error at independent check points "
              f"(condition: {cond}; for a well calibrated σ about 68% of the points fall within 1σ and 95% within 2σ).")
        made.append("fig_uncertainty_coverage")

    # ---- 6. representation-scale compliance (<= 3 conditions so colours stay distinguishable)
    names = list(conds)[:3]
    if names:
        fig, ax = plt.subplots(layout="constrained", figsize=(COL_MM * MM, 55 * MM))
        x = np.arange(len(SCALES))
        w = 0.8 / len(names)
        for k, n in enumerate(names):
            e = np.array(conds[n])
            y = [100 * np.mean(e <= t) for t in SCALES.values()]
            ax.bar(x + (k - (len(names) - 1) / 2) * w, y, w * 0.9, color=SERIES[k], label=n, zorder=3,
                   hatch=["", "////", "...."][k], edgecolor="white", linewidth=0.4)
        ax.set_xticks(x)
        ax.set_xticklabels([f"{k}\n(≤{v:.0f} mm)" for k, v in SCALES.items()])
        ax.set_ylabel("Points within tolerance (%)")
        ax.set_ylim(0, 105)
        ax.grid(axis="x", visible=False)
        ax.legend(loc="upper left", fontsize=6.5)
        _save(fig, out, "fig_scale_compliance", caps,
              "Share of validation points whose 3D error is within the graphic tolerance (0.2 mm on paper) "
              "of each representation scale.")
        made.append("fig_scale_compliance")

    # ---- 7. orientation
    ori = {}
    for r in reps:
        for c in (r.get("gcp") or {}).get("checks") or []:
            if c.get("orientation"):
                ori.setdefault(c["orientation"], []).append(c["err_mm"])
    if len(ori) >= 2:
        order = [o for o in ("verticale", "inclinata", "orizzontale") if o in ori]
        lab = {"verticale": "vertical", "inclinata": "inclined", "orizzontale": "horizontal"}
        fig, ax = plt.subplots(layout="constrained", figsize=(COL_MM * MM, 50 * MM))
        bp = ax.boxplot([ori[o] for o in order], widths=0.5, patch_artist=True,
                        medianprops=dict(color=INK, lw=1.0), whiskerprops=dict(color=INK2, lw=0.6),
                        capprops=dict(color=INK2, lw=0.6), flierprops=dict(marker="o", ms=2.5, mec=INK2, mfc="none"))
        for b in bp["boxes"]:
            b.set(facecolor=SERIES[0] + "33", edgecolor=SERIES[0], lw=0.8)
        ax.set_xticks(range(1, len(order) + 1))
        ax.set_xticklabels([f"{lab[o]}\n(n={len(ori[o])})" for o in order])
        ax.set_ylabel("3D error (mm)")
        ax.set_ylim(bottom=0)
        ax.grid(axis="x", visible=False)
        _save(fig, out, "fig_error_by_orientation", caps,
              f"3D error at check points grouped by surface orientation (condition: {cond}).")
        made.append("fig_error_by_orientation")

    (out / "captions.txt").write_text("\n\n".join(caps), encoding="utf-8")
    return dict(figures=made, font=font, folder=str(out), condition=cond)


def _qualitative(r, out, caps, plt):
    import cv2

    from .imageio_exif import load_image

    files = r.get("files", {})
    dpath = files.get("depth_npy")
    if not dpath or not Path(dpath).exists():
        dpath = str(Path(r["_dir"]) / f"{r['_stem']}_depth_m.npy")
    if not Path(dpath).exists():
        return None
    Z = np.load(dpath)
    if r.get("image_path") and Path(r["image_path"]).exists():
        rgb, _ = load_image(r["image_path"])
    else:
        ov = Path(r["_dir"]) / f"{r['_stem']}_charuco.jpg"
        rgb = cv2.cvtColor(cv2.imread(str(ov)), cv2.COLOR_BGR2RGB)
    H, W = Z.shape
    s = min(1.0, 1600 / max(H, W))
    small = lambda a, nn=False: cv2.resize(a, (int(W * s), int(H * s)),
                                           interpolation=cv2.INTER_NEAREST if nn else cv2.INTER_AREA)
    img = small(rgb)
    Zs = small(Z, nn=True)
    Zm = np.ma.masked_where(Zs <= 0, Zs)
    sig = (r.get("sigma_rel_point") or 0) * 1000
    Sm = np.ma.masked_where(Zs <= 0, Zs * sig)
    pb = {b["board"]: b for b in _chosen_per_board(r)}
    aspect = H / W
    wpanel = (PAGE_MM - 8) / 3
    fig = plt.figure(figsize=(PAGE_MM * MM, (wpanel * aspect + 18) * MM))
    import matplotlib as mpl

    top = 1 - 4 / (wpanel * aspect + 18)
    hfrac = wpanel * aspect / (wpanel * aspect + 18)
    axs = [fig.add_axes([(i * (wpanel + 4)) / PAGE_MM, top - hfrac, wpanel / PAGE_MM, hfrac]) for i in range(3)]
    axs[0].imshow(img)
    iw = img.shape[1]
    for d in r["detections"]:
        x0, y0, x1, y1 = [v * s for v in d["bbox_px"]]
        axs[0].add_patch(mpl.patches.Rectangle((x0, y0), x1 - x0, y1 - y0, fill=False, ec="#eda100", lw=0.8))
        b = pb.get(d["board"])
        if b:
            right = x1 + 0.12 * iw < iw
            axs[0].text(x1 + 4 if right else x0 - 4, y0, f"{b['rmse3d_mm']:.0f} mm", color="white", fontsize=6,
                        va="top", ha="left" if right else "right",
                        bbox=dict(boxstyle="round,pad=0.15", fc="#0b0b0bcc", ec="none"))
    cmap = mpl.colormaps["Blues"].copy()
    cmap.set_bad("#ffffff")
    im1 = axs[1].imshow(Zm, cmap=cmap, vmin=np.percentile(Zs[Zs > 0], 1), vmax=np.percentile(Zs[Zs > 0], 99))
    cmap2 = mpl.colormaps["Oranges"].copy()
    cmap2.set_bad("#ffffff")
    im2 = axs[2].imshow(Sm, cmap=cmap2)
    for a, t in zip(axs, ["(a)", "(b)", "(c)"]):
        a.set_xticks([])
        a.set_yticks([])
        a.grid(False)
        for sp in a.spines.values():
            sp.set_visible(False)
        a.set_title(t, loc="left", pad=2)
    for a, im, lab in [(axs[1], im1, "Depth (m)"), (axs[2], im2, "σ (mm)")]:
        cax = a.inset_axes([0.1, -0.09, 0.8, 0.04])
        cb = fig.colorbar(im, cax=cax, orientation="horizontal")
        cb.set_label(lab, labelpad=1)
        cb.outline.set_linewidth(0.4)
    name = f"fig_qualitative_{r['_stem']}"
    _save(fig, out, name, caps,
          "(a) Input photograph with the detected ChArUco boards and their leave-one-board-out 3D error; "
          "(b) metric depth map after target anchoring; (c) declared 1-sigma uncertainty per point.")
    return name
